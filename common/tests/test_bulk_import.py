"""The spreadsheet import, exercised through the API the panel actually calls."""

import io

import pytest
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.airports.models import Airport
from apps.destinations.tests.factories import DestinationFactory
from apps.hotels.models import Hotel
from apps.users.tests.factories import UserFactory
from apps.visas.models import VisaCountry, VisaType
from common.constants import RoleChoices

pytestmark = pytest.mark.django_db


def sheet(rows: list[list], name="import.xlsx"):
    """A real .xlsx in memory — the same thing a browser would upload."""
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    buffer.name = name
    return buffer


def as_editor():
    client = APIClient()
    client.force_authenticate(user=UserFactory(role=RoleChoices.EDITOR))
    return client


AIRPORT_HEADER = [
    "iata_code", "name_en", "name_ar", "city_en", "city_ar",
    "country_en", "country_ar", "country_code", "is_active",
]
AIRPORT_ROW = ["ZZA", "Test Field", "مطار تجريبي", "Testville", "تستفيل", "Nowhere", "لا مكان", "XX", "yes"]


def test_the_template_is_a_spreadsheet_with_the_columns_to_fill_in():
    response = as_editor().get("/api/v1/airports/import-template/")

    assert response.status_code == 200
    assert "spreadsheetml" in response["Content-Type"]
    workbook = load_workbook(io.BytesIO(response.content))
    headers = [cell.value for cell in workbook.worksheets[0][1]]
    assert "iata_code*" in headers
    # The second sheet is the instructions, so the first can be filled in and
    # sent straight back.
    assert "How to fill this in" in workbook.sheetnames


def test_the_template_is_not_public():
    """It is a tool for whoever maintains the catalogue."""
    assert APIClient().get("/api/v1/airports/import-template/").status_code == 401


def test_a_filled_in_sheet_creates_rows():
    response = as_editor().post(
        "/api/v1/airports/import/",
        {"file": sheet([AIRPORT_HEADER, AIRPORT_ROW])},
        format="multipart",
    )

    assert response.status_code == 200, response.data
    assert response.data["created"] == 1
    assert Airport.objects.get(iata_code="ZZA").city_en == "Testville"


def test_a_row_whose_key_exists_updates_instead_of_duplicating():
    Airport.objects.create(
        iata_code="ZZA", name_en="Old", name_ar="قديم", city_en="Old", city_ar="قديم",
        country_en="Nowhere", country_ar="لا مكان",
    )

    response = as_editor().post(
        "/api/v1/airports/import/",
        {"file": sheet([AIRPORT_HEADER, AIRPORT_ROW])},
        format="multipart",
    )

    assert response.data == {"created": 0, "updated": 1, "errors": [], "preview": [
        {"row": 2, "key": "ZZA", "action": "update"}
    ]}
    assert Airport.objects.filter(iata_code="ZZA").count() == 1
    assert Airport.objects.get(iata_code="ZZA").name_en == "Test Field"


def test_one_bad_row_stops_the_whole_import():
    """A half-applied sheet is worse than a refusal: nobody can tell which half
    went in without reading the rows back one by one."""
    bad = ["ZZB", "", "بلا اسم", "Somewhere", "مكان", "Nowhere", "لا مكان", "XX", "yes"]

    response = as_editor().post(
        "/api/v1/airports/import/",
        {"file": sheet([AIRPORT_HEADER, AIRPORT_ROW, bad])},
        format="multipart",
    )

    assert response.status_code == 400
    assert response.data["created"] == 0
    assert Airport.objects.filter(iata_code__in=["ZZA", "ZZB"]).count() == 0
    error = response.data["errors"][0]
    # Named by the row number the spreadsheet shows, not a zero-based index.
    assert error["row"] == 3
    assert error["column"] == "name_en"


def test_a_dry_run_reports_what_would_happen_and_writes_nothing():
    response = as_editor().post(
        "/api/v1/airports/import/",
        {"file": sheet([AIRPORT_HEADER, AIRPORT_ROW]), "dry_run": "true"},
        format="multipart",
    )

    assert response.status_code == 200
    assert response.data["created"] == 1
    assert response.data["preview"] == [{"row": 2, "key": "ZZA", "action": "create"}]
    assert not Airport.objects.filter(iata_code="ZZA").exists()


def test_a_mistyped_header_is_refused_rather_than_ignored():
    """Silently dropping a column is the failure nobody notices until later."""
    response = as_editor().post(
        "/api/v1/airports/import/",
        {"file": sheet([["iata_code", "nmae_en"], ["ZZA", "Typo Field"]])},
        format="multipart",
    )

    assert response.status_code == 400
    assert "nmae_en" in response.data["errors"][0]["column"]


def test_headers_survive_the_way_people_actually_type_them():
    """Copied out of the template, a header keeps its * — and someone
    retyping it uses capitals and a space."""
    response = as_editor().post(
        "/api/v1/airports/import/",
        {"file": sheet([["IATA_Code*", " Name EN ", "name_ar", "city_en", "city_ar",
                         "country_en", "country_ar"],
                        ["ZZA", "Test Field", "مطار", "City", "مدينة", "Nowhere", "لا مكان"]])},
        format="multipart",
    )

    assert response.status_code == 200, response.data
    assert Airport.objects.filter(iata_code="ZZA").exists()


def test_yes_and_no_are_understood_in_arabic():
    response = as_editor().post(
        "/api/v1/airports/import/",
        {"file": sheet([AIRPORT_HEADER, [*AIRPORT_ROW[:-1], "لا"]])},
        format="multipart",
    )

    assert response.status_code == 200, response.data
    assert Airport.objects.get(iata_code="ZZA").is_active is False


def test_a_cell_that_is_neither_yes_nor_no_says_so():
    response = as_editor().post(
        "/api/v1/airports/import/",
        {"file": sheet([AIRPORT_HEADER, [*AIRPORT_ROW[:-1], "maybe"]])},
        format="multipart",
    )

    assert response.status_code == 400
    assert response.data["errors"][0]["column"] == "is_active"


def test_a_reference_is_given_by_name_rather_than_by_id():
    """Nobody filling in a spreadsheet knows a database id, and asking for one
    is how an import ends up attached to the wrong city."""
    DestinationFactory(name_en="Dubai", name_ar="دبي")

    response = as_editor().post(
        "/api/v1/hotels/import/",
        {
            "file": sheet([
                ["name_en", "name_ar", "destination", "star_rating", "price_per_night_from"],
                ["Test Hotel", "فندق تجريبي", "Dubai", 5, 450],
            ])
        },
        format="multipart",
    )

    assert response.status_code == 200, response.data
    assert Hotel.objects.get(name_en="Test Hotel").destination.name_en == "Dubai"


def test_a_reference_nothing_matches_names_the_row_and_the_column():
    response = as_editor().post(
        "/api/v1/hotels/import/",
        {
            "file": sheet([
                ["name_en", "name_ar", "destination", "star_rating", "price_per_night_from"],
                ["Test Hotel", "فندق", "Atlantis", 5, 450],
            ])
        },
        format="multipart",
    )

    assert response.status_code == 400
    assert response.data["errors"][0]["column"] == "destination"
    assert "Atlantis" in response.data["errors"][0]["message"]


def test_rows_matched_on_two_columns_do_not_overwrite_each_other():
    """"Tourist Visa" exists for every country, so matching on the name alone
    would file Georgia's row over Turkey's."""
    turkey = VisaCountry.objects.create(name_en="Turkey", name_ar="تركيا")
    VisaCountry.objects.create(name_en="Georgia", name_ar="جورجيا")
    VisaType.objects.create(
        country=turkey, name_en="Tourist Visa", name_ar="تأشيرة سياحية",
        price=300, processing_time_days=5,
    )

    response = as_editor().post(
        "/api/v1/visas/import/",
        {
            "file": sheet([
                ["country", "name_en", "name_ar", "price", "processing_time_days"],
                ["Georgia", "Tourist Visa", "تأشيرة سياحية", 250, 3],
            ])
        },
        format="multipart",
    )

    assert response.status_code == 200, response.data
    assert response.data["created"] == 1
    assert VisaType.objects.count() == 2
    assert VisaType.objects.get(country__name_en="Turkey").price == 300


def test_the_public_cannot_import_anything():
    response = APIClient().post(
        "/api/v1/airports/import/",
        {"file": sheet([AIRPORT_HEADER, AIRPORT_ROW])},
        format="multipart",
    )

    assert response.status_code == 401
    assert not Airport.objects.filter(iata_code="ZZA").exists()


def test_a_sheet_with_no_rows_says_so_rather_than_reporting_success():
    response = as_editor().post(
        "/api/v1/airports/import/",
        {"file": sheet([AIRPORT_HEADER])},
        format="multipart",
    )

    assert response.status_code == 400
    assert "no rows" in response.data["errors"][0]["message"]


def test_a_file_that_is_not_a_spreadsheet_is_refused_politely():
    junk = io.BytesIO(b"this is not a workbook")
    junk.name = "notes.xlsx"

    response = as_editor().post("/api/v1/airports/import/", {"file": junk}, format="multipart")

    assert response.status_code == 400
    assert "spreadsheet" in response.data["detail"]


def test_a_csv_exported_from_excel_works_too():
    """Half the people sending a "sheet" send a CSV, and Excel writes a BOM in
    front of the first header."""
    csv_file = io.BytesIO(
        "﻿iata_code,name_en,name_ar,city_en,city_ar,country_en,country_ar\n"
        "ZZA,Test Field,مطار,City,مدينة,Nowhere,لا مكان\n".encode()
    )
    csv_file.name = "airports.csv"

    response = as_editor().post(
        "/api/v1/airports/import/", {"file": csv_file}, format="multipart"
    )

    assert response.status_code == 200, response.data
    assert Airport.objects.filter(iata_code="ZZA").exists()
